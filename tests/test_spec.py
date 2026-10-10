"""声明式模板：YAML 定义 → 渲染 → 校验单元格。"""

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from office_templates.options import Options
from office_templates.render import render_spec
from office_templates.spec import Spec, load_spec

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "specs" / "差旅报销单.yaml"


class LoadSpecTest(unittest.TestCase):
    def test_load_example(self):
        spec = load_spec(EXAMPLE)
        self.assertEqual(spec.filename, "06_差旅报销单.xlsx")
        self.assertEqual([s.name for s in spec.sheets], ["报销明细", "按类型汇总"])

    def test_rejects_unknown_field(self):
        with self.assertRaises(ValueError):
            Spec.from_dict({"filename": "a.xlsx", "sheets": [], "多余的字段": 1})

    def test_rejects_bad_column_type(self):
        with self.assertRaises(ValueError):
            Spec.from_dict(
                {"filename": "a.xlsx", "sheets": [{"name": "S", "columns": [{"label": "X", "type": "随便"}]}]}
            )

    def test_rejects_filename_without_xlsx(self):
        with self.assertRaises(ValueError):
            Spec.from_dict({"filename": "a.xls", "sheets": [{"name": "S", "columns": [{"label": "X"}]}]})

    def test_rejects_missing_file(self):
        with self.assertRaises(ValueError):
            load_spec("不存在的定义.yaml")


class RenderSpecTest(unittest.TestCase):
    def test_render_example(self):
        spec = load_spec(EXAMPLE)
        wb = render_spec(spec, Options(company="某某科技"))

        self.assertEqual(wb.sheetnames, ["报销明细", "按类型汇总"])
        ws = wb["报销明细"]
        self.assertEqual(ws["A1"].value, "某某科技 · 差旅费报销单")
        self.assertEqual([ws.cell(3, j).value for j in range(1, 7)], ["日期", "事由", "类型", "票据张数", "金额", "备注"])
        self.assertEqual(ws["E4"].value, 553)
        self.assertEqual(ws["A4"].value.year, 2026)

        # header_row 3 + rows 30 → 数据 4~33，合计行 34
        self.assertEqual(ws["A34"].value, "合计")
        self.assertEqual(ws["E34"].value, "=SUM(E4:E33)")

        # 汇总表：跨表 SUMIFS 与引用合计行的占比
        s = wb["按类型汇总"]
        self.assertEqual(s["B4"].value, "=SUMIFS(报销明细!$E:$E,报销明细!$C:$C,A4)")
        self.assertEqual(s["C4"].value, "=IFERROR(B4/$B$9,0)")

    def test_dropdown_applied(self):
        spec = load_spec(EXAMPLE)
        ws = render_spec(spec)["报销明细"]
        dv_formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        self.assertIn('"交通,住宿,餐饮,办公用品,其他"', dv_formulas)

    def test_render_and_reopen(self):
        spec = load_spec(EXAMPLE)
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / spec.filename
            render_spec(spec).save(path)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["报销明细", "按类型汇总"])


if __name__ == "__main__":
    unittest.main()
