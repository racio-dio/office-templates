"""冒烟测试：模板能生成、能被重新打开，且参数化确实落到单元格里。"""

import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from office_templates.options import Options
from office_templates.registry import TEMPLATES, build_templates, resolve

EXPECTED_SHEETS = {
    "kaoqin": ["考勤表"],
    "gongzi": ["工资明细", "工资条（打印）"],
    "gantt": ["项目进度"],
    "jizhang": ["流水", "月度汇总", "设置"],
    "kucun": ["库存总览", "出入库记录"],
}


class RegistryTest(unittest.TestCase):
    def test_registry_matches_expected_templates(self):
        self.assertEqual([t.key for t in TEMPLATES], list(EXPECTED_SHEETS))

    def test_resolve_supports_index_and_key(self):
        self.assertEqual([t.key for t in resolve(["1"])], ["kaoqin"])
        self.assertEqual([t.key for t in resolve(["01"])], ["kaoqin"])
        self.assertEqual([t.key for t in resolve(["kucun"])], ["kucun"])
        self.assertEqual([t.key for t in resolve(["1,2", "3"])], ["kaoqin", "gongzi", "gantt"])
        self.assertEqual(len(resolve(None)), len(TEMPLATES))

    def test_resolve_rejects_unknown(self):
        with self.assertRaises(ValueError):
            resolve(["999"])


class BuildTest(unittest.TestCase):
    def test_build_all_and_reopen(self):
        with TemporaryDirectory() as tmp:
            paths = build_templates(TEMPLATES, tmp)
            self.assertEqual(len(paths), len(TEMPLATES))
            for tpl, path in zip(TEMPLATES, paths):
                with self.subTest(template=tpl.key):
                    self.assertTrue(Path(path).is_file())
                    wb = load_workbook(path)
                    self.assertEqual(wb.sheetnames, EXPECTED_SHEETS[tpl.key])


class OptionsTest(unittest.TestCase):
    def test_defaults(self):
        opts = Options()
        self.assertEqual(opts.company, "")
        self.assertEqual(opts.month_tag, "2026-10")
        self.assertEqual(len(opts.staff), 5)

    def test_rejects_bad_month(self):
        with self.assertRaises(ValueError):
            Options(month=13)

    def test_rejects_bad_rate(self):
        with self.assertRaises(ValueError):
            Options(social_rate=1.5)
        with self.assertRaises(ValueError):
            Options(fund_rate=-0.1)

    def test_rejects_empty_staff(self):
        with self.assertRaises(ValueError):
            Options(staff=())

    def test_from_dict_coerces_types(self):
        opts = Options.from_dict(
            {"staff": "甲,乙", "departments": ["研发部"], "start_date": "2026-11-03", "social_rate": "0.08"}
        )
        self.assertEqual(opts.staff, ("甲", "乙"))
        self.assertEqual(opts.departments, ("研发部",))
        self.assertEqual(opts.start_date, date(2026, 11, 3))
        self.assertEqual(opts.social_rate, 0.08)

    def test_rejects_unknown_field(self):
        with self.assertRaises(ValueError):
            Options.from_dict({"company": "X", "不存在的参数": 1})


class ParametrizedBuildTest(unittest.TestCase):
    def test_options_land_in_cells(self):
        opts = Options(
            company="某某科技",
            staff=("赵六", "钱七"),
            departments=("研发部",),
            year=2026,
            month=11,
            social_rate=0.2,
            fund_rate=0.12,
            start_date=date(2026, 11, 3),
        )
        with TemporaryDirectory() as tmp:
            paths = {p.name[:2]: p for p in build_templates(TEMPLATES, tmp, opts)}

            kaoqin = load_workbook(paths["01"])["考勤表"]
            self.assertEqual(kaoqin["A1"].value, "某某科技 · 月度考勤统计表")
            self.assertEqual(kaoqin["B3"].value, 2026)
            self.assertEqual(kaoqin["D3"].value, 11)
            self.assertEqual(kaoqin["B6"].value, "赵六")
            self.assertEqual(kaoqin["B7"].value, "钱七")
            self.assertIsNone(kaoqin["B8"].value, "名单之外不应写入姓名")

            gongzi = load_workbook(paths["02"])["工资明细"]
            self.assertEqual(gongzi["B3"].value, 0.2)
            self.assertEqual(gongzi["D3"].value, 0.12)
            self.assertEqual(gongzi["A6"].value, "A001")
            self.assertEqual(gongzi["B6"].value, "赵六")
            self.assertEqual(gongzi["C6"].value, "研发部")
            self.assertEqual(gongzi["C7"].value, "研发部")

            gantt = load_workbook(paths["03"])["项目进度"]
            self.assertEqual(gantt["B3"].value.date(), date(2026, 11, 3))

            jizhang = load_workbook(paths["04"])["月度汇总"]
            self.assertEqual(jizhang["B3"].value, "2026-11")

    def test_more_staff_than_default_rows(self):
        names = tuple(f"员工{i}" for i in range(1, 41))
        opts = Options(staff=names)
        with TemporaryDirectory() as tmp:
            paths = {p.name[:2]: p for p in build_templates(TEMPLATES, tmp, opts)}
            kaoqin = load_workbook(paths["01"])["考勤表"]
            self.assertEqual(kaoqin["B6"].value, "员工1")
            self.assertEqual(kaoqin["B45"].value, "员工40")


if __name__ == "__main__":
    unittest.main()
