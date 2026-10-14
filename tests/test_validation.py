"""Bad input should fail with a clear ValueError (exit code 2), never a traceback or a broken file."""

import os
import subprocess
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from office_templates.options import Options
from office_templates.render import pick_font, render_spec
from office_templates.spec import Spec

ROOT = Path(__file__).resolve().parents[1]


def spec(sheets=None, **root):
    data = {"filename": "a.xlsx", "sheets": sheets or [{"name": "S", "columns": [{"label": "X"}]}]}
    data.update(root)
    return data


def sheet(**kw):
    data = {"name": "S", "columns": [{"label": "X"}]}
    data.update(kw)
    return data


class SpecValidationTest(unittest.TestCase):
    def bad(self, data):
        with self.assertRaises(ValueError):
            Spec.from_dict(data)

    def test_unknown_placeholder_in_formula(self):
        self.bad(spec([sheet(columns=[{"label": "X", "formula": "=A{row}*2"}])]))

    def test_row_placeholder_not_allowed_in_conditional(self):
        self.bad(spec([sheet(conditionals=[{"range": "A{r}:A{last}", "when": "1"}])]))

    def test_unbalanced_brace(self):
        self.bad(spec([sheet(columns=[{"label": "X", "formula": "=A{r"}])]))

    def test_escaped_braces_still_work(self):
        s = Spec.from_dict(spec([sheet(columns=[{"label": "X", "formula": "=SUM({{1,2}})+A{r}"}])]))
        ws = render_spec(s)["S"]
        self.assertEqual(ws["A4"].value, "=SUM({1,2})+A4")

    def test_duplicate_sheet_names(self):
        self.bad(spec([sheet(name="Items"), sheet(name="items")]))

    def test_sheet_name_too_long(self):
        self.bad(spec([sheet(name="x" * 32)]))

    def test_sheet_name_bad_chars(self):
        for name in ("A/B", "Q?", "[x]", "a:b"):
            with self.subTest(name=name):
                self.bad(spec([sheet(name=name)]))

    def test_header_row_overlapping_title(self):
        self.bad(spec([sheet(header_row=1)]))
        self.bad(spec([sheet(header_row=2)], subtitle="fill the yellow cells"))
        Spec.from_dict(spec([sheet(header_row=2)]))  # no subtitle: row 2 is free

    def test_choices_with_comma(self):
        self.bad(spec([sheet(columns=[{"label": "X", "choices": ["Travel, air", "Meals"]}])]))

    def test_choices_too_long(self):
        self.bad(spec([sheet(columns=[{"label": "X", "choices": ["x" * 30] * 10}])]))

    def test_numeric_strings_are_accepted(self):
        s = Spec.from_dict(spec([sheet(rows="10", header_row="4")]))
        self.assertEqual((s.sheets[0].rows, s.sheets[0].header_row), (10, 4))

    def test_non_numeric_rows(self):
        self.bad(spec([sheet(rows="ten")]))

    def test_wrong_container_types(self):
        self.bad({"filename": "a.xlsx", "sheets": None})
        self.bad(spec([sheet(columns="X")]))
        self.bad(spec([sheet(columns=[{"label": "X", "choices": "a"}])]))

    def test_filename_must_not_be_a_path(self):
        for name in ("../evil.xlsx", "sub/a.xlsx", "sub\\a.xlsx", ".hidden.xlsx"):
            with self.subTest(name=name):
                self.bad(spec(filename=name))


class OptionsCoercionTest(unittest.TestCase):
    def test_string_numbers_from_config(self):
        opts = Options.from_dict({"year": "2027", "month": "3", "fund_rate": "0.05"})
        self.assertEqual((opts.year, opts.month, opts.fund_rate), (2027, 3, 0.05))

    def test_bad_numbers(self):
        for data in ({"month": "March"}, {"social_rate": "abc"}, {"year": -5}, {"month": True}):
            with self.subTest(data=data), self.assertRaises(ValueError):
                Options.from_dict(data)


class FontTest(unittest.TestCase):
    def test_latin_spec_uses_calibri(self):
        s = Spec.from_dict(spec(title="Expense report"))
        self.assertEqual(pick_font(s), "Calibri")
        ws = render_spec(s)["S"]
        self.assertEqual(ws["A1"].font.name, "Calibri")

    def test_cjk_spec_uses_chinese_font(self):
        s = Spec.from_dict(spec(title="报销单"))
        self.assertEqual(pick_font(s), "微软雅黑")


class TotalLabelTest(unittest.TestCase):
    def test_total_label_follows_content_language(self):
        en = render_spec(Spec.from_dict(spec([sheet(rows=2, totals=True)], title="Log")))["S"]
        zh = render_spec(Spec.from_dict(spec([sheet(rows=2, totals=True)], title="台账")))["S"]
        custom = render_spec(Spec.from_dict(spec([sheet(rows=2, totals=True, total_label="Sum")])))["S"]
        self.assertEqual((en["A6"].value, zh["A6"].value, custom["A6"].value), ("Total", "合计", "Sum"))


class PrintSetupTest(unittest.TestCase):
    def test_spec_sheets_print_one_page_wide(self):
        ws = render_spec(Spec.from_dict(spec()))["S"]
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual((ws.page_setup.fitToWidth, ws.page_setup.fitToHeight), (1, 0))


class CliExitCodeTest(unittest.TestCase):
    def run_cli(self, *args):
        env = dict(os.environ, PYTHONPATH=str(ROOT), OFFICE_TEMPLATES_LANG="en")
        return subprocess.run(
            [sys.executable, "-m", "office_templates.cli", *args],
            cwd=ROOT, env=env, capture_output=True, text=True,
        )

    def test_bad_spec_exits_2_without_traceback(self):
        with TemporaryDirectory() as tmp:
            bad = Path(tmp) / "bad.yaml"
            bad.write_text("filename: a.xlsx\nsheets:\n  - name: S\n    columns:\n      - {label: X, formula: '=A{row}'}\n")
            res = self.run_cli("--spec", str(bad), "-o", tmp)
            self.assertEqual(res.returncode, 2)
            self.assertNotIn("Traceback", res.stdout + res.stderr)
            self.assertIn("placeholder", res.stdout + res.stderr)

    def test_english_spec_builds(self):
        with TemporaryDirectory() as tmp:
            res = self.run_cli("--spec", "examples/specs/expense-report.yaml", "-c", "Acme", "-o", tmp)
            self.assertEqual(res.returncode, 0, res.stdout + res.stderr)
            wb = load_workbook(Path(tmp) / "expense-report.xlsx")
            self.assertTrue(wb.worksheets[0]["A1"].value.startswith("Acme"))


if __name__ == "__main__":
    unittest.main()
