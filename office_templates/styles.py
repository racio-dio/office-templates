"""通用样式：配色常量与标题 / 表头 / 正文的快速排版工具。"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

__all__ = [
    "L",
    "B",
    "F",
    "BLUE",
    "HEADER",
    "LIGHT",
    "GRAY",
    "INPUT",
    "INPUT_SOFT",
    "title",
    "header",
    "body",
    "fill_cells",
]

# 配色
BLUE = "1F4E79"  # 主标题底色
HEADER = "2E75B6"  # 表头底色
LIGHT = "DDEBF7"  # 浅蓝，用于强调区域
GRAY = "F2F2F2"  # 正文隔行底色
INPUT = "FFF2CC"  # 需要用户填写的单元格
INPUT_SOFT = "FFF9E5"  # 需要填写的连续区域（比 INPUT 更浅，适合大面积）

thin = Side(style="thin", color="BFBFBF")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
F = "微软雅黑"


def title(ws, text, ncol, sub=None):
    """第一行主标题，可选第二行灰色副标题（填写说明）。"""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, text)
    c.font = Font(name=F, size=16, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        s = ws.cell(2, 1, sub)
        s.font = Font(name=F, size=9, italic=True, color="7F7F7F")
        s.alignment = Alignment(horizontal="center")


def header(ws, row, cols, widths=None):
    """写一行表头，可选同时设置列宽。"""
    for i, h in enumerate(cols, 1):
        c = ws.cell(row, i, h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = B
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[row].height = 24


def body(ws, r1, r2, c1, c2):
    """给一片区域加边框、正文字体、居中与隔行底色。"""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = ws.cell(r, c)
            x.border = B
            x.font = Font(name=F, size=10)
            x.alignment = Alignment(horizontal="center", vertical="center")
            if r % 2 == 0:
                x.fill = PatternFill("solid", fgColor=GRAY)


def fill_cells(ws, r1, r2, cols, color=INPUT):
    """把指定列在 r1~r2 行范围内标成需要填写的颜色。"""
    for r in range(r1, r2 + 1):
        for c in cols:
            ws.cell(r, c).fill = PatternFill("solid", fgColor=color)
