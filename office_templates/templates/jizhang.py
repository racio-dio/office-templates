"""04 个人/家庭记账本。"""

from datetime import date

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..options import Options, with_company
from ..styles import INPUT, L, body, header, title

KEY = "jizhang"
FILENAME = "04_个人家庭记账本.xlsx"
DESC = "分类下拉记账，月度汇总、储蓄率、支出构成饼图自动出"

CATS = ["工资", "副业", "理财", "餐饮", "交通", "购物", "住房", "娱乐", "医疗", "教育", "其他"]
ACCOUNTS = ["微信", "支付宝", "银行卡", "现金", "信用卡"]
ROWS = 500

# （日，收/支，分类，金额，账户，备注）
SAMPLE = [
    (1, "收入", "工资", 12000, "银行卡", "10月工资"),
    (2, "支出", "住房", 3500, "支付宝", "房租"),
    (3, "支出", "餐饮", 68, "微信", "午饭"),
    (5, "支出", "交通", 120, "支付宝", "地铁充值"),
    (6, "支出", "购物", 459, "信用卡", "衣服"),
    (8, "收入", "副业", 800, "微信", "接单"),
]


def build(options: Options) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "流水"
    title(ws, with_company(options.company, "个人/家庭记账本"), 7, "每笔一行，分类用下拉选择；「月度汇总」表自动统计并出图")
    header(ws, 3, ["日期", "收/支", "分类", "金额", "账户", "备注", "月份"], [12, 8, 12, 12, 12, 24, 10])

    lst = wb.create_sheet("设置")
    lst["A1"] = "分类"
    for i, c in enumerate(CATS, 2):
        lst.cell(i, 1, c)
    lst["C1"] = "账户"
    for i, c in enumerate(ACCOUNTS, 2):
        lst.cell(i, 3, c)

    dv1 = DataValidation(type="list", formula1='"收入,支出"')
    dv2 = DataValidation(type="list", formula1="=设置!$A$2:$A$12")
    dv3 = DataValidation(type="list", formula1="=设置!$C$2:$C$6")
    for d in (dv1, dv2, dv3):
        ws.add_data_validation(d)
    dv1.add(f"B4:B{3 + ROWS}")
    dv2.add(f"C4:C{3 + ROWS}")
    dv3.add(f"E4:E{3 + ROWS}")

    for k in range(ROWS):
        r = 4 + k
        if k < len(SAMPLE):
            day, kind, cat, amount, account, note = SAMPLE[k]
            ws.cell(r, 1, date(options.year, options.month, day))
            ws.cell(r, 2, kind)
            ws.cell(r, 3, cat)
            ws.cell(r, 4, amount)
            ws.cell(r, 5, account)
            ws.cell(r, 6, note)
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        ws.cell(r, 4).number_format = "#,##0.00"
        ws.cell(r, 7, f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))')
    body(ws, 4, 40, 1, 7)
    ws.conditional_formatting.add(f"A4:G{3 + ROWS}", FormulaRule(formula=['$B4="收入"'], font=Font(color="548235")))
    ws.freeze_panes = "A4"

    s = wb.create_sheet("月度汇总", 1)
    title(s, "月度汇总", 4)
    s["A3"] = "统计月份"
    s["B3"] = options.month_tag
    s["B3"].fill = PatternFill("solid", fgColor=INPUT)
    s["A4"] = "总收入"
    s["B4"] = '=SUMIFS(流水!D:D,流水!G:G,$B$3,流水!B:B,"收入")'
    s["A5"] = "总支出"
    s["B5"] = '=SUMIFS(流水!D:D,流水!G:G,$B$3,流水!B:B,"支出")'
    s["A6"] = "结余"
    s["B6"] = "=B4-B5"
    s["A7"] = "储蓄率"
    s["B7"] = "=IFERROR(B6/B4,0)"
    s["B7"].number_format = "0.0%"
    for a in ("B4", "B5", "B6"):
        s[a].number_format = "#,##0.00"

    header(s, 9, ["支出分类", "金额", "占比"], [14, 14, 10])
    for i, c in enumerate(CATS[3:], 10):
        s.cell(i, 1, c)
        s.cell(i, 2, f'=SUMIFS(流水!D:D,流水!G:G,$B$3,流水!C:C,A{i},流水!B:B,"支出")')
        s.cell(i, 2).number_format = "#,##0.00"
        s.cell(i, 3, f"=IFERROR(B{i}/$B$5,0)")
        s.cell(i, 3).number_format = "0.0%"
    body(s, 10, 17, 1, 3)

    pie = PieChart()
    pie.title = "支出构成"
    pie.add_data(Reference(s, min_col=2, min_row=9, max_row=17), titles_from_data=True)
    pie.set_categories(Reference(s, min_col=1, min_row=10, max_row=17))
    pie.height = 8
    pie.width = 12
    s.add_chart(pie, "E3")
    return wb
