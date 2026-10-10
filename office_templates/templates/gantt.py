"""03 项目进度甘特图。"""

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ..options import Options, with_company
from ..styles import F, HEADER, INPUT, L, body, header, title

KEY = "gantt"
FILENAME = "03_项目进度甘特图.xlsx"
DESC = "改开始日期和工期，色块自动生成；进度条、今日红线、延期自动提示"

# （任务名，距项目起始日的天数，工期，进度）
TASKS = [
    ("需求调研", 0, 5, 1),
    ("方案设计", 4, 6, 0.8),
    ("开发实现", 9, 15, 0.4),
    ("测试验收", 22, 6, 0),
    ("上线交付", 28, 3, 0),
]
ROWS = 15


def build(options: Options) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "项目进度"
    title(ws, with_company(options.company, "项目进度甘特图"), 50, "修改开始日期和工期，右侧色块自动生成；进度填 0–100%，今天所在列自动标红")
    ws["A3"] = "项目起始日"
    ws["B3"] = options.start_date
    ws["B3"].number_format = "yyyy-mm-dd"
    ws["B3"].fill = PatternFill("solid", fgColor=INPUT)

    cols = ["序号", "任务", "负责人", "开始日期", "工期(天)", "结束日期", "进度", "状态"]
    header(ws, 5, cols, [6, 22, 10, 12, 9, 12, 8, 9])

    staff_count = len(options.staff)
    for k in range(ROWS):
        r = 6 + k
        ws.cell(r, 1, k + 1)
        if k < len(TASKS):
            name, offset, duration, progress = TASKS[k]
            ws.cell(r, 2, name)
            ws.cell(r, 3, options.staff[k % staff_count])
            ws.cell(r, 4, f"=$B$3+{offset}")
            ws.cell(r, 5, duration)
            ws.cell(r, 7, progress)
        ws.cell(r, 4).number_format = "mm-dd"
        ws.cell(r, 6, f'=IF(D{r}="","",D{r}+E{r}-1)')
        ws.cell(r, 6).number_format = "mm-dd"
        ws.cell(r, 7).number_format = "0%"
        ws.cell(r, 8, f'=IF(B{r}="","",IF(G{r}>=1,"已完成",IF(AND(TODAY()>F{r},G{r}<1),"延期",IF(G{r}>0,"进行中","未开始"))))')
    body(ws, 6, 20, 1, 8)

    for i in range(42):
        c = 9 + i
        ws.column_dimensions[L(c)].width = 3.2
        h = ws.cell(5, c, f"=$B$3+{i}")
        h.number_format = "d"
        h.font = Font(name=F, size=8, bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=HEADER)
        h.alignment = Alignment(horizontal="center")
        m = ws.cell(4, c, f'=IF(DAY(I$5)=1,MONTH(I$5)&"月","")' if i == 0 else f'=IF(DAY({L(c)}$5)=1,MONTH({L(c)}$5)&"月","")')
        m.font = Font(name=F, size=8)
        for r in range(6, 21):
            ws.cell(r, c).border = Border(
                left=Side(style="hair", color="D9D9D9"), bottom=Side(style="hair", color="D9D9D9")
            )
    ws.cell(4, 9, '=MONTH(I$5)&"月"')

    g = "I6:AX20"
    ws.conditional_formatting.add(
        g, FormulaRule(formula=["AND(I$5>=$D6,I$5<$D6+ROUND($E6*$G6,0))"], fill=PatternFill("solid", fgColor="2E75B6"))
    )
    ws.conditional_formatting.add(
        g, FormulaRule(formula=["AND(I$5>=$D6,I$5<=$F6)"], fill=PatternFill("solid", fgColor="BDD7EE"))
    )
    ws.conditional_formatting.add(
        "I5:AX20",
        FormulaRule(
            formula=["I$5=TODAY()"],
            border=Border(left=Side(style="medium", color="C00000"), right=Side(style="medium", color="C00000")),
        ),
    )
    ws.conditional_formatting.add(
        "G6:G20", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47")
    )
    ws.conditional_formatting.add(
        "H6:H20", CellIsRule(operator="equal", formula=['"延期"'], font=Font(color="C00000", bold=True))
    )
    ws.freeze_panes = "I6"
    return wb
