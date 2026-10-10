"""05 进销存库存管理。"""

from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..options import Options, with_company
from ..styles import body, header, title

KEY = "kucun"
FILENAME = "05_进销存库存管理.xlsx"
DESC = "登记出入库，当前库存自动汇总，低于安全库存自动标红"

# （编码，名称，规格，单位，期初库存，安全库存）
ITEMS = [
    ("P001", "A4打印纸", "70g", "箱", 20, 10),
    ("P002", "中性笔", "0.5mm", "盒", 50, 20),
    ("P003", "文件夹", "A4", "个", 100, 30),
    ("P004", "订书钉", "24/6", "盒", 30, 10),
]
# （日，编码，类型，数量，经手人）
SAMPLE = [
    (8, "P001", "入库", 5, "张三"),
    (9, "P002", "出库", 40, "李四"),
    (9, "P003", "出库", 20, "王五"),
]
ROWS = 2000


def build(options: Options) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "库存总览"
    title(ws, with_company(options.company, "进销存库存管理表"), 9, "在「出入库记录」登记每笔进出，总览自动算当前库存；低于安全库存自动标红")
    header(ws, 3, ["商品编码", "商品名称", "规格", "单位", "期初库存", "累计入库", "累计出库", "当前库存", "安全库存"], [11, 16, 10, 6, 10, 10, 10, 10, 10])

    rec = wb.create_sheet("出入库记录")
    for k in range(100):
        r = 4 + k
        if k < len(ITEMS):
            for j, v in enumerate(ITEMS[k][:5], 1):
                ws.cell(r, j, v)
            ws.cell(r, 9, ITEMS[k][5])
        ws.cell(r, 6, f'=IF(A{r}="","",SUMIFS(出入库记录!$E:$E,出入库记录!$B:$B,A{r},出入库记录!$D:$D,"入库"))')
        ws.cell(r, 7, f'=IF(A{r}="","",SUMIFS(出入库记录!$E:$E,出入库记录!$B:$B,A{r},出入库记录!$D:$D,"出库"))')
        ws.cell(r, 8, f'=IF(A{r}="","",E{r}+F{r}-G{r})')
    body(ws, 4, 30, 1, 9)
    ws.conditional_formatting.add(
        "A4:I103",
        FormulaRule(
            formula=['AND($A4<>"",$H4<$I4)'],
            fill=PatternFill("solid", fgColor="F8CBAD"),
            font=Font(color="9C0006", bold=True),
        ),
    )
    ws.freeze_panes = "C4"

    title(rec, "出入库记录", 7)
    header(rec, 3, ["日期", "商品编码", "商品名称", "类型", "数量", "经手人", "备注"], [12, 11, 16, 8, 8, 10, 20])
    dv = DataValidation(type="list", formula1='"入库,出库"')
    rec.add_data_validation(dv)
    dv.add(f"D4:D{3 + ROWS}")
    dv2 = DataValidation(type="list", formula1="=库存总览!$A$4:$A$103")
    rec.add_data_validation(dv2)
    dv2.add(f"B4:B{3 + ROWS}")

    staff_count = len(options.staff)
    for k in range(ROWS):
        r = 4 + k
        if k < len(SAMPLE):
            day, code, kind, qty, person = SAMPLE[k]
            rec.cell(r, 1, date(options.year, options.month, day))
            rec.cell(r, 2, code)
            rec.cell(r, 4, kind)
            rec.cell(r, 5, qty)
            rec.cell(r, 6, options.staff[k % staff_count] if k < staff_count else person)
        rec.cell(r, 1).number_format = "yyyy-mm-dd"
        if k < 300:
            rec.cell(r, 3, f'=IFERROR(INDEX(库存总览!$B:$B,MATCH(B{r},库存总览!$A:$A,0)),"")')
    body(rec, 4, 40, 1, 7)
    rec.freeze_panes = "A4"
    return wb
