"""01 月度考勤统计表。"""

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..options import Options, with_company
from ..styles import F, INPUT, L, body, header, title

KEY = "kaoqin"
FILENAME = "01_月度考勤统计表.xlsx"
DESC = "下拉标记 √/迟/假/旷/休，自动统计出勤率，周末自动标红"

MIN_ROWS = 30


def build(options: Options) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "考勤表"
    title(
        ws,
        with_company(options.company, "月度考勤统计表"),
        40,
        "只需修改年份、月份和员工姓名；在日期格里选择 √ 出勤 / 迟 迟到 / 假 请假 / 旷 旷工 / 休 休息，右侧自动统计",
    )
    ws["A3"] = "年份"
    ws["B3"] = options.year
    ws["C3"] = "月份"
    ws["D3"] = options.month
    for a in ("A3", "C3"):
        ws[a].font = Font(name=F, bold=True)
    for a in ("B3", "D3"):
        ws[a].fill = PatternFill("solid", fgColor=INPUT)

    cols = ["序号", "姓名"] + [str(i) for i in range(1, 32)] + ["出勤", "迟到", "请假", "旷工", "休息", "出勤率", "备注"]
    header(ws, 5, cols)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 10
    for i in range(3, 34):
        ws.column_dimensions[L(i)].width = 4.2
        ws.cell(4, i, f'=IFERROR(MID("日一二三四五六",WEEKDAY(DATE($B$3,$D$3,{i - 2})),1),"")')
        ws.cell(4, i).font = Font(name=F, size=8, color="7F7F7F")
        ws.cell(4, i).alignment = Alignment(horizontal="center")
        ws.cell(5, i, f'=IF({i - 2}<=DAY(EOMONTH(DATE($B$3,$D$3,1),0)),{i - 2},"")')
    for i in range(34, 41):
        ws.column_dimensions[L(i)].width = 7
    ws.column_dimensions[L(40)].width = 12

    dv = DataValidation(type="list", formula1='"√,迟,假,旷,休"', allow_blank=True)
    ws.add_data_validation(dv)

    rows = max(MIN_ROWS, len(options.staff))
    last = 5 + rows
    for k in range(rows):
        r = 6 + k
        ws.cell(r, 1, k + 1)
        if k < len(options.staff):
            ws.cell(r, 2, options.staff[k])
        rng = f"C{r}:AG{r}"
        dv.add(rng)
        for j, s in enumerate(["√", "迟", "假", "旷", "休"]):
            ws.cell(r, 34 + j, f'=COUNTIF({rng},"{s}")')
        ws.cell(r, 39, f'=IFERROR((AH{r}+AI{r})/(AH{r}+AI{r}+AJ{r}+AK{r}),"")')
        ws.cell(r, 39).number_format = "0.0%"
    body(ws, 6, last, 1, 40)

    for s, color in [("迟", "FFE699"), ("假", "BDD7EE"), ("旷", "F8CBAD"), ("休", "E2EFDA")]:
        ws.conditional_formatting.add(
            f"C6:AG{last}",
            CellIsRule(operator="equal", formula=[f'"{s}"'], fill=PatternFill("solid", fgColor=color)),
        )
    ws.conditional_formatting.add(
        f"C4:AG{last}", FormulaRule(formula=['OR(C$4="六",C$4="日")'], font=Font(color="C00000"))
    )
    ws.freeze_panes = "C6"
    return wb
